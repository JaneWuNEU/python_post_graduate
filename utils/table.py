# -*- coding: utf-8 -*-
"""
Created on Sat Feb 25 21:16:25 2017

@author: User
"""
from docx import Document

def createTable():
    doc = Document()
    table = doc.add_table(rows=9,cols = 9)
    cell = table.cell(0,1)
    cell.text = "work"
    
    doc.save("F:/test.docx")
createTable()